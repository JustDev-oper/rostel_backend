import csv
from io import StringIO, BytesIO
from docx import Document

def make_csv(items: list[dict]) -> str:
    if not items: return ''
    output = StringIO(); writer = csv.DictWriter(output, fieldnames=items[0].keys(), extrasaction='ignore'); writer.writeheader(); writer.writerows(items)
    return output.getvalue()

def make_proposal_docx(proposal: dict, deal: dict | None) -> bytes:
    doc = Document(); doc.add_heading('Коммерческое предложение', 0)
    doc.add_paragraph(f"№ {proposal['number']}"); doc.add_heading(proposal['title'], 1)
    doc.add_paragraph(proposal['content'] or 'Описание предложения отсутствует.')
    table = doc.add_table(rows=0, cols=2); table.style = 'Table Grid'
    for key, value in [('Стоимость', f"{proposal['total_price']:,.2f} ₽"), ('Действительно до', proposal.get('valid_until') or 'Не указано'), ('Статус', proposal['status'])]:
        cells = table.add_row().cells; cells[0].text = key; cells[1].text = str(value)
    if deal: doc.add_paragraph(f"Сделка: {deal['name']}")
    doc.add_paragraph('Документ сформирован автоматически. Тестовые данные отмечены в API.')
    out = BytesIO(); doc.save(out); return out.getvalue()
